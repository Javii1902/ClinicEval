import math
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook

from ..db import get_connection
from .assessment_config import ASSESSMENT_ITEMS, MONTH_OPTIONS
from ..settings import get_compliance_threshold, set_compliance_threshold
from ..validators import parse_test_counts, classify_answer
from .excel_exports import (
    sanitize_sheet_name,
    calculate_percentage,
    format_percentage,
    write_dashboard_sheet,
    write_clinic_sheet,
    build_detail_from_dashboard_row,
)
from .pdf_exports import export_dashboard_pdf
from .scroll_utils import bind_mousewheel_scrolling, bind_bounded_mousewheel_scrolling
from .widgets import MultiSelectPopup, CanvasTooltip

QUESTIONS = [item["text"] for item in ASSESSMENT_ITEMS]


class Dashboard(ttk.Frame):
    CLINIC_HEADER_LABELS = {
        "clinic_name": "Clinic Name",
        "region": "Region",
        "specialty": "Specialty",
        "compliance": "Compliance Percentage",
        "yes": "Yes",
        "no": "No",
        "stl": "STL",
        "na": "N/A",
    }

    QUESTION_HEADER_LABELS = {
        "question": "Question",
        "compliance": "Compliance Percentage",
        "yes": "Yes",
        "no": "No",
        "stl": "STL",
        "na": "N/A",
    }

    TEST_CHART_COLORS = [
        "#4E79A7",
        "#F28E2B",
        "#E15759",
        "#76B7B2",
        "#59A14F",
        "#EDC948",
        "#B07AA1",
        "#FF9DA7",
        "#9C755F",
        "#BAB0AC",
        "#1F77B4",
        "#FF7F0E",
    ]

    PIE_COLORS = {
        "Yes": "#59A14F",
        "No": "#E15759",
        "STL": "#F28E2B",
        "N/A": "#76B7B2",
    }

    MONTH_LABEL_BY_NUMBER = {
        option.split(" - ")[0]: option for option in MONTH_OPTIONS if option != "All"
    }

    MAX_BAR_CHART_HEIGHT = 480

    def __init__(self, parent, app):
        super().__init__(parent)
        self.app = app

        self.total_assessments_var = tk.StringVar(value="0")
        self.total_clinics_var = tk.StringVar(value="0")
        self.overall_compliance_var = tk.StringVar(value="N/A")
        self.compliant_clinics_var = tk.StringVar(value="0")
        self.non_compliant_clinics_var = tk.StringVar(value="0")

        self.year_var = tk.StringVar(value="All")
        self.month_var = tk.StringVar(value="All")
        self.compliance_threshold_var = tk.StringVar(value=str(get_compliance_threshold()))

        self.available_clinics = []
        self.available_regions = []
        self.available_specialties = []

        self.selected_clinics = []
        self.selected_regions = []
        self.selected_specialties = []

        self.clinics_filter_is_all = True
        self.regions_filter_is_all = True
        self.specialties_filter_is_all = True

        self.year_combobox = None
        self.month_combobox = None
        self.clinic_filter_button = None
        self.region_filter_button = None
        self.specialty_filter_button = None

        self.clinic_tree = None
        self.question_tree = None

        self.bar_chart_canvas = None
        self.bar_chart_legend_frame = None
        self.bar_chart_tooltip = None
        self.pie_chart_canvas = None
        self.pie_chart_tooltip = None
        self.timeline_canvas = None
        self.timeline_tooltip = None

        self.last_dataset = None
        self.last_timeline_dataset = None
        self.chart_redraw_job = None

        self.clinic_sort_column = "clinic_name"
        self.clinic_sort_descending = False
        self.question_sort_column = "question"
        self.question_sort_descending = False

        self.page_canvas = None
        self.page_scrollbar = None
        self.page_container = None
        self.page_window = None

        self.build_ui()
        self.after(150, self.refresh_data)

    def build_ui(self):
        self.page_canvas = tk.Canvas(self, highlightthickness=0)
        self.page_scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.page_canvas.yview)
        self.page_container = ttk.Frame(self.page_canvas)

        self.page_container.bind(
            "<Configure>",
            lambda e: self.page_canvas.configure(scrollregion=self.page_canvas.bbox("all"))
        )

        self.page_window = self.page_canvas.create_window((0, 0), window=self.page_container, anchor="nw")
        self.page_canvas.configure(yscrollcommand=self.page_scrollbar.set)

        self.page_canvas.pack(side="left", fill="both", expand=True)
        self.page_scrollbar.pack(side="right", fill="y")

        self.page_canvas.bind("<Configure>", self.on_page_canvas_configure)
        bind_mousewheel_scrolling(self.page_canvas)

        outer = ttk.Frame(self.page_container, padding=12)
        outer.pack(fill="both", expand=True)

        header_frame = ttk.Frame(outer)
        header_frame.pack(fill="x", pady=(0, 10))

        ttk.Label(header_frame, text="Dashboard", font=("Segoe UI", 14, "bold")).pack(side="left")
        ttk.Button(header_frame, text="Export to Excel", command=self.export_dashboard_workbook).pack(side="right", padx=(6, 0))
        ttk.Button(header_frame, text="Export to PDF", command=self.export_dashboard_pdf).pack(side="right", padx=(6, 0))
        ttk.Button(header_frame, text="Refresh", command=self.refresh_data).pack(side="right")

        filter_frame = ttk.Frame(outer)
        filter_frame.pack(fill="x", pady=(0, 12))

        ttk.Label(filter_frame, text="Year").grid(row=0, column=0, padx=(0, 6), sticky="w")
        self.year_combobox = ttk.Combobox(filter_frame, textvariable=self.year_var, width=10, state="readonly")
        self.year_combobox.grid(row=0, column=1, padx=(0, 12), sticky="w")
        self.year_combobox.bind("<<ComboboxSelected>>", lambda e: self.on_date_filter_changed())

        ttk.Label(filter_frame, text="Month").grid(row=0, column=2, padx=(0, 6), sticky="w")
        self.month_combobox = ttk.Combobox(
            filter_frame,
            textvariable=self.month_var,
            width=16,
            state="readonly",
            values=["All"],
        )
        self.month_combobox.grid(row=0, column=3, padx=(0, 12), sticky="w")
        self.month_combobox.bind("<<ComboboxSelected>>", lambda e: self.on_date_filter_changed())

        ttk.Label(filter_frame, text="Regions").grid(row=0, column=4, padx=(0, 6), sticky="w")
        self.region_filter_button = ttk.Button(filter_frame, text="All Regions", command=self.open_region_filter_popup, width=20)
        self.region_filter_button.grid(row=0, column=5, padx=(0, 12), sticky="w")

        ttk.Label(filter_frame, text="Specialties").grid(row=0, column=6, padx=(0, 6), sticky="w")
        self.specialty_filter_button = ttk.Button(filter_frame, text="All Specialties", command=self.open_specialty_filter_popup, width=20)
        self.specialty_filter_button.grid(row=0, column=7, padx=(0, 12), sticky="w")

        ttk.Label(filter_frame, text="Clinics").grid(row=0, column=8, padx=(0, 6), sticky="w")
        self.clinic_filter_button = ttk.Button(filter_frame, text="All Clinics", command=self.open_clinic_filter_popup, width=20)
        self.clinic_filter_button.grid(row=0, column=9, padx=(0, 12), sticky="w")

        ttk.Button(filter_frame, text="Clear Filters", command=self.clear_filters).grid(row=0, column=10, sticky="w")

        ttk.Label(filter_frame, text="Compliant at ≥").grid(row=0, column=11, padx=(16, 4), sticky="w")
        compliance_threshold_entry = ttk.Entry(filter_frame, textvariable=self.compliance_threshold_var, width=5)
        compliance_threshold_entry.grid(row=0, column=12, sticky="w")
        compliance_threshold_entry.bind("<Return>", lambda e: self.apply_compliance_threshold())
        ttk.Label(filter_frame, text="%").grid(row=0, column=13, padx=(2, 6), sticky="w")
        ttk.Button(filter_frame, text="Apply", command=self.apply_compliance_threshold).grid(row=0, column=14, sticky="w")

        cards_frame = ttk.Frame(outer)
        cards_frame.pack(fill="x", pady=(0, 4))

        self.build_card(cards_frame, "Total Assessments", self.total_assessments_var, 0)
        self.build_card(cards_frame, "Total Clinics", self.total_clinics_var, 1)
        self.build_card(cards_frame, "Overall Compliance", self.overall_compliance_var, 2)
        self.build_card(cards_frame, "Compliant Clinics", self.compliant_clinics_var, 3)
        self.build_card(cards_frame, "Non-Compliant Clinics", self.non_compliant_clinics_var, 4)

        for i in range(5):
            cards_frame.columnconfigure(i, weight=1)

        ttk.Label(
            outer,
            text="Total Assessments/Total Clinics count every assessment/distinct clinic matching the "
                 "current filters. Overall Compliance, Compliant Clinics, and Non-Compliant Clinics use "
                 "only each clinic's most recent matching assessment - a clinic with more than one "
                 "assessment in the filtered range only counts once, using its latest result. A clinic "
                 "whose latest assessment has no Yes/No/STL answers at all isn't counted as either "
                 "compliant or non-compliant, since there's nothing to grade.",
            font=("Segoe UI", 8),
            foreground="#666666",
            wraplength=1100,
            justify="left",
        ).pack(anchor="w", pady=(0, 12))

        timeline_section = ttk.LabelFrame(outer, text="Compliance Timeline", padding=10)
        timeline_section.pack(fill="both", expand=True, pady=(0, 12))

        ttk.Label(
            timeline_section,
            text="Each point is each included clinic's Nth assessment averaged together - point 1 is "
                 "every clinic's 1st assessment, point 2 is every clinic's 2nd, and so on - so this "
                 "tracks progress by visit number rather than by calendar date, honoring the current "
                 "Region/Specialty/Clinic and Year/Month filters. Hover a point to see every clinic's "
                 "individual compliance for that visit.",
            font=("Segoe UI", 8),
            foreground="#666666",
            wraplength=1100,
            justify="left",
        ).pack(anchor="w", pady=(0, 6))

        self.timeline_canvas = tk.Canvas(
            timeline_section,
            height=260,
            background="white",
            highlightthickness=1,
            highlightbackground="#d9d9d9",
        )
        self.timeline_canvas.pack(fill="both", expand=True)
        self.timeline_canvas.bind("<Configure>", lambda e: self.schedule_chart_redraw())
        self.timeline_tooltip = CanvasTooltip(self.timeline_canvas)

        charts_section = ttk.Frame(outer)
        charts_section.pack(fill="x", pady=(0, 12))
        charts_section.columnconfigure(0, weight=3)
        charts_section.columnconfigure(1, weight=2)

        bar_section = ttk.LabelFrame(charts_section, text="Tests Performed Bar Chart", padding=10)
        bar_section.grid(row=0, column=0, sticky="nsew", padx=(0, 8))

        self.bar_chart_canvas = tk.Canvas(
            bar_section,
            height=320,
            background="white",
            highlightthickness=1,
            highlightbackground="#d9d9d9",
        )
        self.bar_chart_canvas.pack(fill="both", expand=True)
        self.bar_chart_canvas.bind("<Configure>", lambda e: self.schedule_chart_redraw())
        self.bar_chart_tooltip = CanvasTooltip(self.bar_chart_canvas)

        self.bar_chart_legend_frame = ttk.Frame(bar_section)
        self.bar_chart_legend_frame.pack(fill="x", pady=(8, 0))

        pie_section = ttk.LabelFrame(charts_section, text="Answer Distribution Pie Chart", padding=10)
        pie_section.grid(row=0, column=1, sticky="nsew")

        self.pie_chart_canvas = tk.Canvas(
            pie_section,
            height=320,
            background="white",
            highlightthickness=1,
            highlightbackground="#d9d9d9",
        )
        self.pie_chart_canvas.pack(fill="both", expand=True)
        self.pie_chart_canvas.bind("<Configure>", lambda e: self.schedule_chart_redraw())
        self.pie_chart_tooltip = CanvasTooltip(self.pie_chart_canvas)

        clinic_section = ttk.Frame(outer)
        clinic_section.pack(fill="both", expand=True, pady=(0, 12))

        ttk.Label(clinic_section, text="Clinic Compliance Table", font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(0, 6))

        clinic_tree_frame = ttk.Frame(clinic_section)
        clinic_tree_frame.pack(fill="both", expand=True)

        clinic_columns = ("clinic_name", "region", "specialty", "compliance", "yes", "no", "stl", "na")
        self.clinic_tree = ttk.Treeview(clinic_tree_frame, columns=clinic_columns, show="headings", height=8)
        self.clinic_tree.pack(side="left", fill="both", expand=True)

        clinic_scrollbar = ttk.Scrollbar(clinic_tree_frame, orient="vertical", command=self.clinic_tree.yview)
        clinic_scrollbar.pack(side="right", fill="y")
        self.clinic_tree.configure(yscrollcommand=clinic_scrollbar.set)
        bind_bounded_mousewheel_scrolling(self.clinic_tree, self.page_canvas)

        self.clinic_tree.heading("clinic_name", command=lambda: self.sort_treeview(self.clinic_tree, "clinic_name", is_numeric=False))
        self.clinic_tree.heading("region", command=lambda: self.sort_treeview(self.clinic_tree, "region", is_numeric=False))
        self.clinic_tree.heading("specialty", command=lambda: self.sort_treeview(self.clinic_tree, "specialty", is_numeric=False))
        self.clinic_tree.heading("compliance", command=lambda: self.sort_treeview(self.clinic_tree, "compliance", is_numeric=True, is_percentage=True))
        self.clinic_tree.heading("yes", command=lambda: self.sort_treeview(self.clinic_tree, "yes", is_numeric=True))
        self.clinic_tree.heading("no", command=lambda: self.sort_treeview(self.clinic_tree, "no", is_numeric=True))
        self.clinic_tree.heading("stl", command=lambda: self.sort_treeview(self.clinic_tree, "stl", is_numeric=True))
        self.clinic_tree.heading("na", command=lambda: self.sort_treeview(self.clinic_tree, "na", is_numeric=True))

        self.clinic_tree.column("clinic_name", anchor="w", width=220, stretch=True)
        self.clinic_tree.column("region", anchor="w", width=120, stretch=True)
        self.clinic_tree.column("specialty", anchor="w", width=180, stretch=True)
        self.clinic_tree.column("compliance", anchor="e", width=150, stretch=False)
        self.clinic_tree.column("yes", anchor="e", width=70, stretch=False)
        self.clinic_tree.column("no", anchor="e", width=70, stretch=False)
        self.clinic_tree.column("stl", anchor="e", width=70, stretch=False)
        self.clinic_tree.column("na", anchor="e", width=70, stretch=False)

        question_section = ttk.Frame(outer)
        question_section.pack(fill="both", expand=True)

        ttk.Label(question_section, text="Question Compliance Table", font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(0, 6))

        question_tree_frame = ttk.Frame(question_section)
        question_tree_frame.pack(fill="both", expand=True)

        question_columns = ("question", "compliance", "yes", "no", "stl", "na")
        self.question_tree = ttk.Treeview(question_tree_frame, columns=question_columns, show="headings", height=12)
        self.question_tree.pack(side="left", fill="both", expand=True)

        question_scrollbar = ttk.Scrollbar(question_tree_frame, orient="vertical", command=self.question_tree.yview)
        question_scrollbar.pack(side="right", fill="y")
        self.question_tree.configure(yscrollcommand=question_scrollbar.set)
        bind_bounded_mousewheel_scrolling(self.question_tree, self.page_canvas)

        self.question_tree.heading("question", command=lambda: self.sort_treeview(self.question_tree, "question", is_numeric=False))
        self.question_tree.heading("compliance", command=lambda: self.sort_treeview(self.question_tree, "compliance", is_numeric=True, is_percentage=True))
        self.question_tree.heading("yes", command=lambda: self.sort_treeview(self.question_tree, "yes", is_numeric=True))
        self.question_tree.heading("no", command=lambda: self.sort_treeview(self.question_tree, "no", is_numeric=True))
        self.question_tree.heading("stl", command=lambda: self.sort_treeview(self.question_tree, "stl", is_numeric=True))
        self.question_tree.heading("na", command=lambda: self.sort_treeview(self.question_tree, "na", is_numeric=True))

        self.question_tree.column("question", anchor="w", width=700, stretch=True)
        self.question_tree.column("compliance", anchor="e", width=150, stretch=False)
        self.question_tree.column("yes", anchor="e", width=70, stretch=False)
        self.question_tree.column("no", anchor="e", width=70, stretch=False)
        self.question_tree.column("stl", anchor="e", width=70, stretch=False)
        self.question_tree.column("na", anchor="e", width=70, stretch=False)

        self.update_sort_indicators()

    def on_page_canvas_configure(self, event):
        self.page_canvas.itemconfigure(self.page_window, width=event.width)

    def build_card(self, parent, title, value_var, column):
        card = ttk.Frame(parent, padding=12, relief="ridge")
        card.grid(row=0, column=column, padx=6, sticky="nsew")
        ttk.Label(card, text=title, font=("Segoe UI", 9, "bold")).pack(anchor="w")
        ttk.Label(card, textvariable=value_var, font=("Segoe UI", 16, "bold")).pack(anchor="w", pady=(8, 0))

    def refresh_data(self):
        self.load_filter_values()
        self.refresh_filter_options()
        self.load_dashboard()

    def on_date_filter_changed(self):
        self.load_filter_values()
        self.refresh_filter_options()
        self.load_dashboard()

    def clear_filters(self):
        self.year_var.set("All")
        self.month_var.set("All")
        self.selected_clinics = []
        self.selected_regions = []
        self.selected_specialties = []
        self.clinics_filter_is_all = True
        self.regions_filter_is_all = True
        self.specialties_filter_is_all = True
        self.load_filter_values()
        self.refresh_filter_options()
        self.load_dashboard()

    def focus_clinic(self, clinic_name):
        # Jumped to from the Clinic tab's "View Dashboard" button - resets
        # every other filter (Year/Month/Region/Specialty back to "All") so
        # the requested clinic's full history is visible, then narrows the
        # clinic filter down to just that one clinic.
        self.year_var.set("All")
        self.month_var.set("All")
        self.selected_regions = []
        self.selected_specialties = []
        self.regions_filter_is_all = True
        self.specialties_filter_is_all = True
        self.selected_clinics = [clinic_name] if clinic_name else []
        self.clinics_filter_is_all = False

        self.load_filter_values()
        self.refresh_filter_options()
        self.load_dashboard()

    def apply_compliance_threshold(self):
        raw_value = self.compliance_threshold_var.get().strip()
        try:
            value = float(raw_value)
        except ValueError:
            messagebox.showerror("Validation Error", "Compliance threshold must be a number between 0 and 100.")
            self.compliance_threshold_var.set(str(get_compliance_threshold()))
            return

        if value < 0 or value > 100:
            messagebox.showerror("Validation Error", "Compliance threshold must be between 0 and 100.")
            self.compliance_threshold_var.set(str(get_compliance_threshold()))
            return

        set_compliance_threshold(value)
        self.load_dashboard()

    def get_available_years(self):
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT DISTINCT substr(assessment_date, 1, 4) AS year_part
                FROM assessments
                WHERE assessment_date IS NOT NULL AND assessment_date != ''
                ORDER BY year_part DESC
            """)
            return [row[0] for row in cursor.fetchall() if row[0]]

    def get_available_months(self, year_filter):
        sql = """
            SELECT DISTINCT substr(assessment_date, 6, 2) AS month_part
            FROM assessments
            WHERE assessment_date IS NOT NULL AND assessment_date != ''
        """
        params = []

        if year_filter and year_filter != "All":
            sql += " AND substr(assessment_date, 1, 4) = ? "
            params.append(year_filter)

        sql += " ORDER BY month_part"

        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            return [row[0] for row in cursor.fetchall() if row[0]]

    def load_filter_values(self):
        try:
            years = self.get_available_years()
            year_values = ["All"] + years

            current_year = self.year_var.get()
            if current_year not in year_values:
                current_year = "All"
                self.year_var.set("All")

            self.year_combobox["values"] = year_values

            available_months = self.get_available_months(current_year)
            month_values = ["All"] + [self.MONTH_LABEL_BY_NUMBER.get(m, m) for m in available_months]
            self.month_combobox["values"] = month_values

            if self.month_var.get() not in month_values:
                self.month_var.set("All")

        except Exception as e:
            messagebox.showerror("Dashboard Error", f"Could not load dashboard filters.\n\n{e}")

    def build_common_date_sql(self):
        sql = " FROM assessments WHERE clinic_name IS NOT NULL AND TRIM(clinic_name) != '' "
        params = []

        selected_year = self.year_var.get().strip()
        selected_month = self.month_var.get().strip()

        if selected_year and selected_year != "All":
            sql += " AND substr(assessment_date, 1, 4) = ? "
            params.append(selected_year)

        if selected_month and selected_month != "All":
            month_number = selected_month.split(" - ")[0]
            sql += " AND substr(assessment_date, 6, 2) = ? "
            params.append(month_number)

        return sql, params

    def get_distinct_regions(self):
        sql_base, params = self.build_common_date_sql()
        sql = f"""
            SELECT DISTINCT region
            {sql_base}
            AND region IS NOT NULL AND TRIM(region) != ''
            ORDER BY region
        """
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            return [row[0] for row in cursor.fetchall() if row[0]]

    def get_distinct_specialties(self):
        sql_base, params = self.build_common_date_sql()

        if self.selected_regions:
            placeholders = ",".join("?" for _ in self.selected_regions)
            sql_base += f" AND region IN ({placeholders}) "
            params.extend(self.selected_regions)

        sql = f"""
            SELECT DISTINCT specialty
            {sql_base}
            AND specialty IS NOT NULL AND TRIM(specialty) != ''
            ORDER BY specialty
        """
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            return [row[0] for row in cursor.fetchall() if row[0]]

    def get_distinct_clinics(self):
        sql_base, params = self.build_common_date_sql()

        if self.selected_regions:
            placeholders = ",".join("?" for _ in self.selected_regions)
            sql_base += f" AND region IN ({placeholders}) "
            params.extend(self.selected_regions)

        if self.selected_specialties:
            placeholders = ",".join("?" for _ in self.selected_specialties)
            sql_base += f" AND specialty IN ({placeholders}) "
            params.extend(self.selected_specialties)

        sql = f"""
            SELECT DISTINCT clinic_name
            {sql_base}
            AND clinic_name IS NOT NULL AND TRIM(clinic_name) != ''
            ORDER BY clinic_name
        """
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            return [row[0] for row in cursor.fetchall() if row[0]]

    def reconcile_selection(self, selected, available, is_all):
        if is_all:
            return list(available), True

        narrowed = [item for item in selected if item in available]
        if not narrowed and available:
            return list(available), True

        return narrowed, False

    def refresh_filter_options(self):
        try:
            regions = self.get_distinct_regions()
            self.available_regions = regions
            self.selected_regions, self.regions_filter_is_all = self.reconcile_selection(
                self.selected_regions, regions, self.regions_filter_is_all
            )

            specialties = self.get_distinct_specialties()
            self.available_specialties = specialties
            self.selected_specialties, self.specialties_filter_is_all = self.reconcile_selection(
                self.selected_specialties, specialties, self.specialties_filter_is_all
            )

            clinics = self.get_distinct_clinics()
            self.available_clinics = clinics
            self.selected_clinics, self.clinics_filter_is_all = self.reconcile_selection(
                self.selected_clinics, clinics, self.clinics_filter_is_all
            )

            self.update_filter_button_labels()

        except Exception as e:
            messagebox.showerror("Dashboard Error", f"Could not refresh filters.\n\n{e}")

    def update_filter_button_labels(self):
        self.clinic_filter_button.config(text=self.build_filter_label(self.selected_clinics, self.available_clinics, "All Clinics"))
        self.region_filter_button.config(text=self.build_filter_label(self.selected_regions, self.available_regions, "All Regions"))
        self.specialty_filter_button.config(text=self.build_filter_label(self.selected_specialties, self.available_specialties, "All Specialties"))

    def build_filter_label(self, selected, available, all_label):
        if not available:
            return all_label
        if len(selected) == 0:
            return "None Selected"
        if len(selected) == len(available):
            return all_label
        if len(selected) == 1:
            return selected[0]
        return f"{len(selected)} Selected"

    def open_clinic_filter_popup(self):
        MultiSelectPopup(self, "Select Clinics", self.available_clinics, self.selected_clinics, self.apply_clinic_filter)

    def open_region_filter_popup(self):
        MultiSelectPopup(self, "Select Regions", self.available_regions, self.selected_regions, self.apply_region_filter)

    def open_specialty_filter_popup(self):
        MultiSelectPopup(self, "Select Specialties", self.available_specialties, self.selected_specialties, self.apply_specialty_filter)

    def apply_clinic_filter(self, selected_values):
        self.selected_clinics = list(selected_values)
        self.clinics_filter_is_all = set(self.selected_clinics) == set(self.available_clinics)
        self.load_dashboard()
        self.update_filter_button_labels()

    def apply_region_filter(self, selected_values):
        self.selected_regions = list(selected_values)
        self.regions_filter_is_all = set(self.selected_regions) == set(self.available_regions)
        self.refresh_filter_options()
        self.load_dashboard()

    def apply_specialty_filter(self, selected_values):
        self.selected_specialties = list(selected_values)
        self.specialties_filter_is_all = set(self.selected_specialties) == set(self.available_specialties)
        self.refresh_filter_options()
        self.load_dashboard()

    def clear_tree(self, tree):
        for item in tree.get_children():
            tree.delete(item)

    def clear_frame(self, frame):
        for widget in frame.winfo_children():
            widget.destroy()

    def parse_tests_performed(self, value):
        return parse_test_counts(value)

    def update_sort_indicators(self):
        for col, label in self.CLINIC_HEADER_LABELS.items():
            arrow = ""
            if col == self.clinic_sort_column:
                arrow = " ↓" if self.clinic_sort_descending else " ↑"
            self.clinic_tree.heading(col, text=label + arrow)

        for col, label in self.QUESTION_HEADER_LABELS.items():
            arrow = ""
            if col != "question" and col == self.question_sort_column:
                arrow = " ↓" if self.question_sort_descending else " ↑"
            self.question_tree.heading(col, text=label + arrow)

    def sort_treeview(self, tree, column, is_numeric=False, is_percentage=False, initial_load=False):
        items = list(tree.get_children(""))

        if tree == self.clinic_tree:
            current_column = self.clinic_sort_column
            if initial_load:
                descending = self.clinic_sort_descending
            elif current_column == column:
                descending = not self.clinic_sort_descending
            else:
                descending = is_numeric or is_percentage

            self.clinic_sort_column = column
            self.clinic_sort_descending = descending

            def convert(value):
                text = (value or "").strip()
                if is_percentage:
                    if text.upper() == "N/A" or not text:
                        return float("-inf")
                    return float(text.replace("%", ""))
                if is_numeric:
                    try:
                        return float(text)
                    except ValueError:
                        return float("-inf")
                return text.lower()

            items.sort(key=lambda item_id: convert(tree.set(item_id, column)), reverse=descending)

        else:
            if column == "question":
                self.question_sort_column = "question"
                self.question_sort_descending = False

                def question_key(item_id):
                    try:
                        return int(str(item_id).replace("q", ""))
                    except ValueError:
                        return 0

                items.sort(key=question_key, reverse=False)
            else:
                current_column = self.question_sort_column
                if initial_load:
                    descending = self.question_sort_descending
                elif current_column == column:
                    descending = not self.question_sort_descending
                else:
                    descending = is_numeric or is_percentage

                self.question_sort_column = column
                self.question_sort_descending = descending

                def convert(value):
                    text = (value or "").strip()
                    if is_percentage:
                        if text.upper() == "N/A" or not text:
                            return float("-inf")
                        return float(text.replace("%", ""))
                    if is_numeric:
                        try:
                            return float(text)
                        except ValueError:
                            return float("-inf")
                    return text.lower()

                items.sort(key=lambda item_id: convert(tree.set(item_id, column)), reverse=descending)

        for index, item_id in enumerate(items):
            tree.move(item_id, "", index)

        self.update_sort_indicators()

    def build_assessment_filter_where(self):
        # Shared by get_filtered_rows() (needs every matching row, for the
        # true assessment count) and get_total_assessment_count() (needs
        # only the count) so the two never drift out of sync on what
        # counts as "matching the current filters".
        selected_year = self.year_var.get().strip()
        selected_month = self.month_var.get().strip()

        where = " WHERE clinic_name IS NOT NULL AND TRIM(clinic_name) != '' "
        params = []

        if selected_year and selected_year != "All":
            where += " AND substr(assessment_date, 1, 4) = ? "
            params.append(selected_year)

        if selected_month and selected_month != "All":
            month_number = selected_month.split(" - ")[0]
            where += " AND substr(assessment_date, 6, 2) = ? "
            params.append(month_number)

        if self.available_regions and self.selected_regions:
            placeholders = ",".join("?" for _ in self.selected_regions)
            where += f" AND region IN ({placeholders}) "
            params.extend(self.selected_regions)
        elif self.available_regions and not self.selected_regions:
            return None, None

        if self.available_specialties and self.selected_specialties:
            placeholders = ",".join("?" for _ in self.selected_specialties)
            where += f" AND specialty IN ({placeholders}) "
            params.extend(self.selected_specialties)
        elif self.available_specialties and not self.selected_specialties:
            return None, None

        if self.available_clinics and self.selected_clinics:
            placeholders = ",".join("?" for _ in self.selected_clinics)
            where += f" AND clinic_name IN ({placeholders}) "
            params.extend(self.selected_clinics)
        elif self.available_clinics and not self.selected_clinics:
            return None, None

        return where, params

    def get_total_assessment_count(self):
        where, params = self.build_assessment_filter_where()
        if where is None:
            return 0

        sql = f"SELECT COUNT(*) FROM assessments {where}"

        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            row = cursor.fetchone()

        return row[0] if row else 0

    def get_filtered_rows(self):
        where, params = self.build_assessment_filter_where()
        if where is None:
            return []

        sql = f"""
            SELECT
                id, clinic_name, region, specialty, assessment_date, tests_performed, notes,
                q1_answer, q1_template_comment, q1_custom_comment,
                q2_answer, q2_template_comment, q2_custom_comment,
                q3_answer, q3_template_comment, q3_custom_comment,
                q4_answer, q4_template_comment, q4_custom_comment,
                q5_answer, q5_template_comment, q5_custom_comment,
                q6_answer, q6_template_comment, q6_custom_comment,
                q7_answer, q7_template_comment, q7_custom_comment,
                q8_answer, q8_template_comment, q8_custom_comment,
                q9_answer, q9_template_comment, q9_custom_comment,
                q10_answer, q10_template_comment, q10_custom_comment,
                q11_answer, q11_template_comment, q11_custom_comment,
                q12_answer, q12_template_comment, q12_custom_comment,
                q13_answer, q13_template_comment, q13_custom_comment,
                q14_answer, q14_template_comment, q14_custom_comment,
                q15_answer, q15_template_comment, q15_custom_comment
            FROM assessments
            {where}
            ORDER BY clinic_name ASC, assessment_date DESC, id DESC
        """

        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            rows = cursor.fetchall()

        latest_by_clinic = {}
        for row in rows:
            clinic_name = (row[1] or "").strip()
            if clinic_name and clinic_name not in latest_by_clinic:
                latest_by_clinic[clinic_name] = row

        return list(latest_by_clinic.values())

    def build_dashboard_dataset(self):
        rows = self.get_filtered_rows()
        compliance_threshold = get_compliance_threshold()

        overall_yes = overall_no = overall_stl = overall_na = 0
        compliant_clinics = 0
        non_compliant_clinics = 0
        clinic_rows = []
        test_counts = {}
        test_clinic_counts = {}
        question_stats = [{"yes": 0, "no": 0, "stl": 0, "na": 0} for _ in ASSESSMENT_ITEMS]

        for row in rows:
            yes_count = no_count = stl_count = na_count = 0

            for idx, _item in enumerate(ASSESSMENT_ITEMS):
                answer = (row[7 + (idx * 3)] or "").strip()
                classification = classify_answer(answer)

                if classification == "yes":
                    yes_count += 1
                    overall_yes += 1
                    question_stats[idx]["yes"] += 1
                elif classification == "no":
                    no_count += 1
                    overall_no += 1
                    question_stats[idx]["no"] += 1
                elif classification == "stl":
                    stl_count += 1
                    overall_stl += 1
                    question_stats[idx]["stl"] += 1
                else:
                    na_count += 1
                    overall_na += 1
                    question_stats[idx]["na"] += 1

            compliance = calculate_percentage(yes_count, no_count, stl_count)
            # A clinic with zero Yes/No/STL answers has no compliance to
            # grade at all (calculate_percentage returns None) - it isn't a
            # compliance *failure*, just ungraded, so it shouldn't be forced
            # into "non-compliant" regardless of how low the threshold is.
            if compliance is not None:
                if compliance >= compliance_threshold:
                    compliant_clinics += 1
                else:
                    non_compliant_clinics += 1

            clinic_rows.append((
                row[1] or "",
                row[2] or "",
                row[3] or "",
                compliance,
                yes_count,
                no_count,
                stl_count,
                na_count,
                row,
            ))

            clinic_name = row[1] or ""
            for test_name, count in self.parse_tests_performed(row[5] or "").items():
                test_counts[test_name] = test_counts.get(test_name, 0) + count
                # get_filtered_rows() already dedups to one (latest) row per
                # clinic, so each clinic contributes at most once per test -
                # a plain assignment, not an accumulation, is correct here.
                test_clinic_counts.setdefault(test_name, {})[clinic_name] = count

        clinic_rows.sort(key=lambda x: (x[0] or "").strip().lower())

        return {
            "rows": rows,
            "total_assessment_count": self.get_total_assessment_count(),
            "clinic_rows": clinic_rows,
            "overall_yes": overall_yes,
            "overall_no": overall_no,
            "overall_stl": overall_stl,
            "overall_na": overall_na,
            "compliant_clinics": compliant_clinics,
            "non_compliant_clinics": non_compliant_clinics,
            "test_counts": test_counts,
            "test_clinic_counts": test_clinic_counts,
            "question_stats": question_stats,
        }

    def build_timeline_dataset(self):
        clinics = self.selected_clinics
        if not clinics:
            return {"points": []}

        placeholders = ",".join("?" for _ in clinics)
        answer_columns = ", ".join(f"q{i}_answer" for i in range(1, 16))
        sql = f"""
            SELECT clinic_name, assessment_date, id, {answer_columns}
            FROM assessments
            WHERE clinic_name IN ({placeholders})
        """
        params = list(clinics)

        selected_year = self.year_var.get().strip()
        selected_month = self.month_var.get().strip()

        if selected_year and selected_year != "All":
            sql += " AND substr(assessment_date, 1, 4) = ? "
            params.append(selected_year)

        if selected_month and selected_month != "All":
            month_number = selected_month.split(" - ")[0]
            sql += " AND substr(assessment_date, 6, 2) = ? "
            params.append(month_number)

        sql += " ORDER BY clinic_name ASC, assessment_date ASC, id ASC "

        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql, params)
            rows = cursor.fetchall()

        by_clinic = {}
        for row in rows:
            clinic_name = (row[0] or "").strip()
            if not clinic_name:
                continue
            by_clinic.setdefault(clinic_name, []).append(row)

        # Bucket by each clinic's OWN assessment sequence (1st, 2nd, 3rd...)
        # rather than by calendar time, so the chart tracks how clinics
        # trend across their own visits regardless of what date each
        # clinic happened to be assessed on.
        sequence_buckets = {}
        for clinic_name, clinic_rows in by_clinic.items():
            for seq_index, row in enumerate(clinic_rows, start=1):
                assessment_date = row[1]
                answers = row[3:18]

                yes_count = no_count = stl_count = 0
                for answer in answers:
                    answer = (answer or "").strip()
                    if answer == "Yes":
                        yes_count += 1
                    elif answer == "No":
                        no_count += 1
                    elif answer == "STL":
                        stl_count += 1

                compliance = calculate_percentage(yes_count, no_count, stl_count)

                bucket = sequence_buckets.setdefault(seq_index, [])
                bucket.append({
                    "clinic_name": clinic_name,
                    "date": assessment_date,
                    "compliance": compliance,
                })

        points = []
        for seq_index in sorted(sequence_buckets.keys()):
            clinic_entries = sequence_buckets[seq_index]
            values = [c["compliance"] for c in clinic_entries if c["compliance"] is not None]
            average = sum(values) / len(values) if values else None

            points.append({
                "sequence": seq_index,
                "label": f"Assessment {seq_index}",
                "average": average,
                "clinic_count": len(clinic_entries),
                "clinics": sorted(clinic_entries, key=lambda c: (c["clinic_name"] or "").strip().lower()),
            })

        return {"points": points}

    def compliance_color(self, pct):
        pct = max(0.0, min(100.0, pct))
        red = (211, 47, 47)
        yellow = (245, 166, 35)
        green = (67, 160, 71)

        if pct <= 50:
            ratio = pct / 50
            start, end = red, yellow
        else:
            ratio = (pct - 50) / 50
            start, end = yellow, green

        rgb = tuple(round(start[i] + (end[i] - start[i]) * ratio) for i in range(3))
        return "#{:02x}{:02x}{:02x}".format(*rgb)

    def draw_timeline_chart(self, dataset):
        self.timeline_canvas.delete("all")
        self.timeline_tooltip.set_regions([])

        points = [p for p in dataset.get("points", []) if p["average"] is not None]

        if not points:
            self.timeline_canvas.create_text(
                20, 25,
                anchor="w",
                text="No timeline data available for the current filters.",
                font=("Segoe UI", 10),
                fill="#555555",
            )
            return

        self.timeline_canvas.update_idletasks()
        width = max(self.timeline_canvas.winfo_width(), 620)
        height = max(self.timeline_canvas.winfo_height(), 260)

        left_pad = 55
        right_pad = 30
        top_pad = 30
        bottom_pad = 40

        chart_width = width - left_pad - right_pad
        chart_height = height - top_pad - bottom_pad
        count = len(points)

        def x_for(index):
            if count == 1:
                return left_pad + chart_width / 2
            return left_pad + chart_width * (index / (count - 1))

        def y_for(value):
            return top_pad + chart_height * (1 - (value / 100))

        for pct in (0, 20, 40, 60, 80, 100):
            y = y_for(pct)
            self.timeline_canvas.create_line(left_pad, y, width - right_pad, y, fill="#eeeeee")
            self.timeline_canvas.create_text(
                left_pad - 8, y,
                anchor="e",
                text=f"{pct}%",
                font=("Segoe UI", 8),
                fill="#444444",
            )

        self.timeline_canvas.create_line(left_pad, top_pad, left_pad, height - bottom_pad, fill="#666666")
        self.timeline_canvas.create_line(
            left_pad, height - bottom_pad, width - right_pad, height - bottom_pad, fill="#666666"
        )

        coords = [(x_for(idx), y_for(point["average"])) for idx, point in enumerate(points)]

        if len(coords) > 1:
            flat = [value for xy in coords for value in xy]
            self.timeline_canvas.create_line(*flat, fill="#999999", width=2)

        regions = []
        for idx, (point, (x, y)) in enumerate(zip(points, coords)):
            # Centered text on the first/last point sits right at the plot's
            # edge and spills past the canvas boundary (e.g. "Assessment 4"
            # getting clipped) - anchoring it inward instead keeps every
            # label fully on-screen without needing extra padding.
            if idx == 0 and count > 1:
                label_anchor = "w"
            elif idx == count - 1 and count > 1:
                label_anchor = "e"
            else:
                label_anchor = "center"

            color = self.compliance_color(point["average"])
            self.timeline_canvas.create_oval(x - 6, y - 6, x + 6, y + 6, fill=color, outline="#333333")

            self.timeline_canvas.create_text(
                x, y - 14,
                text=f"{point['average']:.1f}%",
                anchor=label_anchor,
                font=("Segoe UI", 9, "bold"),
                fill="black",
            )

            self.timeline_canvas.create_text(
                x, height - bottom_pad + 14,
                text=point["label"],
                anchor=label_anchor,
                font=("Segoe UI", 9),
                fill="#333333",
            )

            lines = [
                point["label"],
                f"Average Compliance: {point['average']:.1f}%",
                f"Clinics: {point['clinic_count']}",
                "",
            ]
            for clinic in point["clinics"]:
                compliance_text = (
                    f"{clinic['compliance']:.1f}%" if clinic["compliance"] is not None else "N/A"
                )
                lines.append(f"  {clinic['clinic_name']}: {compliance_text} ({clinic['date']})")

            regions.append({
                "type": "circle",
                "circle": (x, y, 9),
                "id": idx,
                "lines": lines,
            })

        self.timeline_tooltip.set_regions(regions)

    def build_test_color_map(self, test_names):
        return {
            test_name: self.TEST_CHART_COLORS[index % len(self.TEST_CHART_COLORS)]
            for index, test_name in enumerate(test_names)
        }

    def draw_bar_chart(self, dataset, redraw_only=False):
        self.bar_chart_canvas.delete("all")
        self.bar_chart_tooltip.set_regions([])
        if not redraw_only:
            self.clear_frame(self.bar_chart_legend_frame)

        test_counts = dataset.get("test_counts", {})
        if not test_counts:
            self.bar_chart_canvas.create_text(
                20, 25,
                anchor="w",
                text="No test data available.",
                font=("Segoe UI", 10),
                fill="#555555",
            )
            return

        test_names = list(test_counts.keys())
        max_count = max(test_counts.values()) if test_counts else 1

        # A fixed indexed palette instead of randomly generated colors: this
        # chart can now redraw on every window resize (see schedule_chart_
        # redraw), and re-rolling random colors on each of those redraws was
        # pure wasted work for a result that's already deterministic (same
        # seed every time) - a direct lookup is both simpler and free.
        colors = self.build_test_color_map(test_names)

        self.bar_chart_canvas.update_idletasks()
        width = max(self.bar_chart_canvas.winfo_width(), 620)
        height = max(self.bar_chart_canvas.winfo_height(), 320)

        left_pad = 60
        right_pad = 30
        top_pad = 30
        bottom_pad = 30
        number_col_width = 35
        chart_start_x = left_pad + number_col_width
        chart_width = width - chart_start_x - right_pad
        row_height = 34
        bar_height = 20

        # Cap how tall this canvas tries to grow: the bar and pie charts share
        # a grid row, and letting this grow unbounded with the number of test
        # names (previously up to 1000+px for large test lists) dragged the
        # pie chart's canvas up to match, shrinking its actual circle to a
        # speck relative to all the blank space - looking like it vanished.
        # Past the cap, rows get more compact instead of the canvas growing
        # further. Row height still has a legibility floor, though, so with
        # enough tests the compacted total can exceed the cap - in that case
        # every test must still get its own visible row, so needed_height is
        # always recomputed from the final row_height rather than clamped to
        # the cap outright (clamping it unconditionally previously cut off
        # whichever tests didn't fit inside the under-sized canvas).
        needed_height = top_pad + bottom_pad + (len(test_names) * row_height)
        if needed_height > self.MAX_BAR_CHART_HEIGHT:
            available_rows_height = self.MAX_BAR_CHART_HEIGHT - top_pad - bottom_pad
            row_height = max(16, available_rows_height / len(test_names))
            bar_height = max(8, row_height - 12)
            needed_height = top_pad + bottom_pad + (len(test_names) * row_height)

        if height < needed_height:
            self.bar_chart_canvas.config(height=needed_height)
            height = needed_height

        self.bar_chart_canvas.create_line(chart_start_x, top_pad - 5, chart_start_x, height - bottom_pad,
                                          fill="#666666")
        self.bar_chart_canvas.create_line(chart_start_x, height - bottom_pad, width - right_pad, height - bottom_pad,
                                          fill="#666666")

        tick_count = 5
        for i in range(tick_count + 1):
            tick_value = int(round((max_count / tick_count) * i))
            x = chart_start_x + (chart_width * (i / tick_count))
            self.bar_chart_canvas.create_line(x, top_pad - 2, x, height - bottom_pad, fill="#eeeeee")
            self.bar_chart_canvas.create_text(
                x,
                height - bottom_pad + 14,
                text=str(tick_value),
                font=("Segoe UI", 8),
                fill="#444444",
            )

        bar_regions = []
        for index, test_name in enumerate(test_names, start=1):
            value = test_counts[test_name]
            color = colors[test_name]

            y = top_pad + ((index - 1) * row_height)
            y0 = y
            y1 = y0 + bar_height

            bar_width = chart_width * (value / max_count if max_count else 0)
            x0 = chart_start_x
            x1 = x0 + bar_width

            self.bar_chart_canvas.create_text(
                left_pad,
                y0 + (bar_height / 2),
                text=str(index),
                anchor="center",
                font=("Segoe UI", 9, "bold"),
                fill="black",
            )

            self.bar_chart_canvas.create_rectangle(
                x0, y0, x1, y1,
                fill=color,
                outline=color,
            )

            self.bar_chart_canvas.create_text(
                x1 + 8,
                y0 + (bar_height / 2),
                text=str(value),
                anchor="w",
                font=("Segoe UI", 9),
                fill="black",
            )

            clinic_breakdown = dataset.get("test_clinic_counts", {}).get(test_name, {})
            sorted_breakdown = sorted(
                clinic_breakdown.items(),
                key=lambda item: (-item[1], item[0].lower()),
            )
            tooltip_lines = [test_name, f"Total machines: {value}", ""]
            tooltip_lines.extend(
                f"  {clinic_name}: {clinic_count}" for clinic_name, clinic_count in sorted_breakdown
            )

            # Hover region spans the full row (not just the bar) so it's easy
            # to trigger the tooltip even when the bar itself is short.
            bar_regions.append({
                "type": "rect",
                "bbox": (chart_start_x, y0, width - right_pad, y1),
                "id": index,
                "lines": tooltip_lines,
            })

        self.bar_chart_tooltip.set_regions(bar_regions)

        if redraw_only:
            # The legend's layout is a fixed 2-column grid that never
            # depends on canvas size, so a resize-triggered redraw has
            # nothing to gain from destroying and recreating every legend
            # widget (up to ~90 for a 30-test chart) - only rebuild it when
            # the underlying data actually changed.
            return

        for index, test_name in enumerate(test_names, start=1):
            row = (index - 1) // 2
            col = (index - 1) % 2

            item_frame = ttk.Frame(self.bar_chart_legend_frame)
            item_frame.grid(row=row, column=col, sticky="w", padx=(0, 18), pady=2)

            swatch = tk.Canvas(item_frame, width=14, height=14, highlightthickness=0)
            swatch.pack(side="left", padx=(0, 6))
            swatch.create_rectangle(1, 1, 13, 13, fill=colors[test_name], outline=colors[test_name])

            ttk.Label(item_frame, text=f"{index} - {test_name} ({test_counts[test_name]})").pack(side="left")

    def draw_pie_chart(self, dataset):
        self.pie_chart_canvas.delete("all")
        self.pie_chart_tooltip.set_regions([])

        values = {
            "Compliant": dataset.get("compliant_clinics", 0),
            "Non-Compliant": dataset.get("non_compliant_clinics", 0),
        }

        colors = {
            "Compliant": "#4F81BD",
            "Non-Compliant": "#C0504D",
        }

        total = sum(values.values())
        if total <= 0:
            self.pie_chart_canvas.create_text(
                20, 25,
                anchor="w",
                text="No compliance data available.",
                font=("Segoe UI", 10),
                fill="#555555",
            )
            return

        self.pie_chart_canvas.update_idletasks()
        width = max(self.pie_chart_canvas.winfo_width(), 320)
        height = max(self.pie_chart_canvas.winfo_height(), 280)

        diameter = min(width * 0.55, height * 0.72)
        x0 = 20
        y0 = (height - diameter) / 2
        x1 = x0 + diameter
        y1 = y0 + diameter
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2

        wedge_regions = []
        start = 90
        for label, value in values.items():
            if value <= 0:
                continue

            # Tkinter's create_arc silently fails to draw anything when extent
            # is exactly 360 (a single category holding 100% of the data) -
            # it collapses to an invisible sliver instead of a full circle.
            # Clamping just under 360 keeps it rendering as a full circle.
            extent = min((value / total) * 360, 359.9)
            self.pie_chart_canvas.create_arc(
                x0, y0, x1, y1,
                start=start,
                extent=extent,
                fill=colors[label],
                outline="white",
                width=2,
            )

            mid_angle = start + (extent / 2)
            label_radius = diameter * 0.33
            lx = cx + math.cos(math.radians(-mid_angle)) * label_radius
            ly = cy + math.sin(math.radians(-mid_angle)) * label_radius
            self.pie_chart_canvas.create_text(
                lx, ly,
                text=f"Count,\n{value}",
                font=("Segoe UI", 9),
                justify="center",
                fill="black",
            )

            wedge_regions.append({
                "type": "wedge",
                "wedge": (cx, cy, diameter / 2, start, extent),
                "id": label,
                "lines": [label, f"Count: {value}"],
            })

            start += extent

        self.pie_chart_tooltip.set_regions(wedge_regions)

        legend_x = x1 + 18
        legend_y = y0 + 35

        row = 0
        for label, value in values.items():
            self.pie_chart_canvas.create_rectangle(
                legend_x,
                legend_y + (row * 28),
                legend_x + 10,
                legend_y + 10 + (row * 28),
                fill=colors[label],
                outline=colors[label],
            )
            self.pie_chart_canvas.create_text(
                legend_x + 16,
                legend_y + 5 + (row * 28),
                anchor="w",
                text=label,
                font=("Segoe UI", 9),
                fill="black",
            )
            row += 1

    def apply_current_sort(self):
        self.sort_treeview(
            self.clinic_tree,
            self.clinic_sort_column,
            is_numeric=self.clinic_sort_column in {"yes", "no", "stl", "na", "compliance"},
            is_percentage=self.clinic_sort_column == "compliance",
            initial_load=True,
        )
        self.sort_treeview(
            self.question_tree,
            self.question_sort_column,
            is_numeric=self.question_sort_column in {"yes", "no", "stl", "na", "compliance"},
            is_percentage=self.question_sort_column == "compliance",
            initial_load=True,
        )

    def load_dashboard(self):
        dataset = self.build_dashboard_dataset()

        total_assessments = dataset["total_assessment_count"]
        total_clinics = len(dataset["clinic_rows"])
        overall_compliance = calculate_percentage(
            dataset["overall_yes"],
            dataset["overall_no"],
            dataset["overall_stl"],
        )

        self.total_assessments_var.set(str(total_assessments))
        self.total_clinics_var.set(str(total_clinics))
        self.overall_compliance_var.set(format_percentage(overall_compliance))
        self.compliant_clinics_var.set(str(dataset["compliant_clinics"]))
        self.non_compliant_clinics_var.set(str(dataset["non_compliant_clinics"]))

        self.clear_tree(self.clinic_tree)
        for clinic_name, region, specialty, compliance, yes_count, no_count, stl_count, na_count, _source_row in dataset["clinic_rows"]:
            self.clinic_tree.insert(
                "",
                "end",
                values=(
                    clinic_name,
                    region,
                    specialty,
                    format_percentage(compliance),
                    yes_count,
                    no_count,
                    stl_count,
                    na_count,
                )
            )

        self.clear_tree(self.question_tree)
        for idx, item in enumerate(ASSESSMENT_ITEMS):
            stats = dataset["question_stats"][idx]
            compliance = calculate_percentage(stats["yes"], stats["no"], stats["stl"])
            self.question_tree.insert(
                "",
                "end",
                iid=f"q{idx + 1}",
                values=(
                    item["text"],
                    format_percentage(compliance),
                    stats["yes"],
                    stats["no"],
                    stats["stl"],
                    stats["na"],
                )
            )

        self.apply_current_sort()

        self.last_dataset = dataset
        self.last_timeline_dataset = self.build_timeline_dataset()
        self.draw_bar_chart(self.last_dataset)
        self.draw_pie_chart(self.last_dataset)
        self.draw_timeline_chart(self.last_timeline_dataset)

    def schedule_chart_redraw(self):
        if self.chart_redraw_job is not None:
            self.after_cancel(self.chart_redraw_job)
        self.chart_redraw_job = self.after(150, self.redraw_charts_only)

    def redraw_charts_only(self):
        self.chart_redraw_job = None
        if self.last_dataset is not None:
            self.draw_bar_chart(self.last_dataset, redraw_only=True)
            self.draw_pie_chart(self.last_dataset)
        if self.last_timeline_dataset is not None:
            self.draw_timeline_chart(self.last_timeline_dataset)

    def export_dashboard_workbook(self):
        try:
            dataset = self.build_dashboard_dataset()
            if not dataset["rows"]:
                messagebox.showwarning("No Data", "There is no filtered dashboard data to export.")
                return

            file_path = filedialog.asksaveasfilename(
                title="Save Dashboard Workbook",
                defaultextension=".xlsx",
                filetypes=[("Excel Workbook", "*.xlsx")],
                initialfile="dashboard_export.xlsx",
            )
            if not file_path:
                return

            wb = Workbook()
            ws = wb.active

            write_dashboard_sheet(
                ws=ws,
                dataset=dataset,
                questions=QUESTIONS,
                year_value=self.year_var.get(),
                month_value=self.month_var.get(),
                selected_clinics=self.selected_clinics,
                selected_regions=self.selected_regions,
                selected_specialties=self.selected_specialties,
                timeline_dataset=self.build_timeline_dataset(),
            )

            for clinic_name, _region, _specialty, _compliance, _yes, _no, _stl, _na, source_row in dataset["clinic_rows"]:
                detail = build_detail_from_dashboard_row(source_row, ASSESSMENT_ITEMS)
                sheet_name = sanitize_sheet_name(clinic_name or "Clinic", "Clinic")
                clinic_ws = wb.create_sheet(title=sheet_name)
                write_clinic_sheet(clinic_ws, detail, ASSESSMENT_ITEMS, sheet_title=sheet_name)

            wb.save(file_path)
            messagebox.showinfo("Export Complete", f"Dashboard workbook exported successfully.\n\n{file_path}")

        except Exception as e:
            messagebox.showerror("Export Error", f"Could not export dashboard workbook.\n\n{e}")

    def export_dashboard_pdf(self):
        try:
            dataset = self.build_dashboard_dataset()
            if not dataset["rows"]:
                messagebox.showwarning("No Data", "There is no filtered dashboard data to export.")
                return

            file_path = filedialog.asksaveasfilename(
                title="Save Dashboard PDF",
                defaultextension=".pdf",
                filetypes=[("PDF Document", "*.pdf")],
                initialfile="dashboard_export.pdf",
            )
            if not file_path:
                return

            export_dashboard_pdf(
                file_path=file_path,
                dataset=dataset,
                questions=QUESTIONS,
                year_value=self.year_var.get(),
                month_value=self.month_var.get(),
                selected_clinics=self.selected_clinics,
                selected_regions=self.selected_regions,
                selected_specialties=self.selected_specialties,
                timeline_dataset=self.build_timeline_dataset(),
            )

            messagebox.showinfo("Export Complete", f"Dashboard PDF exported successfully.\n\n{file_path}")

        except Exception as e:
            messagebox.showerror("Export Error", f"Could not export dashboard PDF.\n\n{e}")


